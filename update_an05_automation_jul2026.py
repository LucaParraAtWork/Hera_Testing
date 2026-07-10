"""
update_an05_automation_jul2026.py
Credit the new AN_05 "Delete slot entirely" automation added to
test_weekly_ui.py (2026-07-03) on the corresponding Excel row and recolor it
from YELLOW (gap) to GREEN (automated).

  Nom_AN_05 -> test_weekly_ui.py AN_05 (_delete_slot_entirely): opens a NEW
  slot, deletes every trailer row (not just one, unlike UI_05), Saves, and
  confirms the slot returns to a truly empty grid cell (no status-new/
  rejected, no content — same heuristic _add_slot_ui/_snapshot_slots use).

This closes out all three Additional Nomination automation gaps flagged on
2026-07-02 (Nom_AN_02, Nom_AN_03, Nom_AN_05).

Not yet run live — same caveat as AN_02/AN_03.

Run:
    python update_an05_automation_jul2026.py
"""
import openpyxl
from openpyxl.styles import Alignment, PatternFill

PATH   = r"C:\Users\LucaParraGabarron\Desktop\ViryaH2\Hera_Testing\Hera_UAT_Test_Plan.xlsx"
SCRIPT = "test_weekly_ui.py"

WRAP  = Alignment(wrap_text=True, vertical="top")
GREEN = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")

PATCHES = {
    "Nom_AN_05": {
        "Extra description ": (
            "[AUTOMATED 2026-07-03] Covered by test_weekly_ui.py — AN_05 "
            "(_delete_slot_entirely): opens a NEW slot, deletes every trailer row "
            "(not just one, unlike UI_05 / Nom_AN_04), Saves, and confirms the slot "
            "returns to a truly empty grid cell (no status-new/rejected, no content — "
            "same emptiness heuristic used by _add_slot_ui/UI_12). NOT yet run live — "
            "selector/save behaviour may need adjustment on the first real run."
        ),
        "Created via script": SCRIPT,
    },
}

wb = openpyxl.load_workbook(PATH)
ws = wb.active

hdr = [str(c.value or "") for c in ws[1]]
COL = {h: i + 1 for i, h in enumerate(hdr)}   # 1-indexed for openpyxl, no .strip()

updated = []
for row in ws.iter_rows(min_row=2):
    tc_id_cell = row[COL["Test case ID"] - 1]
    tc_id = str(tc_id_cell.value or "").strip()
    if tc_id in PATCHES:
        patch = PATCHES[tc_id]
        for col_name, new_val in patch.items():
            col_idx = COL.get(col_name)
            if col_idx:
                cell = row[col_idx - 1]
                cell.value = new_val
                cell.alignment = WRAP
        for cell in row:
            cell.fill = GREEN
        updated.append(tc_id)
        print(f"  Updated {tc_id} (recolored GREEN)")

if not updated:
    print("  No matching rows found — check Test case ID column.")
else:
    wb.save(PATH)
    print(f"\nSaved {PATH}")
    print(f"Updated {len(updated)} rows: {updated}")
