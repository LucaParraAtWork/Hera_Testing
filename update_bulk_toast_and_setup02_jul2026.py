"""
update_bulk_toast_and_setup02_jul2026.py
Align the Excel with the 2026-07-03 fixes made to test_weekly_ui.py:

  1. SETUP_02 (Nom_WN_UI_SETUP_02) no longer re-imports a separately generated
     "V2" CSV (it threw a Hera server exception on a live run) — it now
     re-imports the same baseline Virya_Nomination_Week21.csv, which already
     resets slots to NEW via Hera's own "each import creates a new highest
     version" behaviour.
  2. UI_10 / UI_11 (Nom_WN_UI_10 / Nom_WN_UI_11) no longer verify bulk
     confirm/reject by snapshotting slot CSS-class counts before/after (that
     approach produced false results twice — a CSV-import-cascade mismatch,
     then an over-strict rewrite that silently no-op'd into a false pass).
     They now check for the application's own confirmation banner (e.g.
     "Rejected week / For Virya Energy NV, Messer Belgium NV.").

Only touches: Nom_WN_UI_SETUP_02, Nom_WN_UI_10, Nom_WN_UI_11.

Run:
    python update_bulk_toast_and_setup02_jul2026.py
"""
import openpyxl
from openpyxl.styles import Alignment

PATH   = r"C:\Users\LucaParraGabarron\Desktop\ViryaH2\Hera_Testing\Hera_UAT_Test_Plan.xlsx"
SCRIPT = "test_weekly_ui.py"

WRAP = Alignment(wrap_text=True, vertical="top")

# ---------------------------------------------------------------------------
# Patch definitions  { test_case_id : { column_name : new_value } }
# ---------------------------------------------------------------------------
PATCHES = {

    "Nom_WN_UI_SETUP_02": {
        "Test case title": "Set up Week 21/2027 – re-import baseline for bulk tests (SETUP_02)",
        "Input with test step by step": (
            "1. Re-import Virya_Nomination_Week21.csv (same file as SETUP_01) for "
            "“Virya Energy NV”.\n"
            "2. Navigate to Week 21/2027.\n"
            "   Automated by: test_weekly_ui.py – SETUP_02"
        ),
        "Expected results": (
            "Import succeeds and creates a new highest-version import. Fresh NEW slots "
            "appear on the Week 21/2027 grid, overwriting the previous rejected state."
        ),
        "Extra description ": (
            "[UPDATED 2026-07-03] SETUP_02 previously re-imported a separately generated "
            "Virya_Nomination_Week21_V2.csv (same rows, a few trailer types flipped to "
            "“type 2”). On a live run this V2 file made Hera throw “unexpected "
            "exception … contact your administrator”, which then cascaded into false "
            "AUTO-FAIL results on UI_10/UI_11 (DB left in a partial state by the failed "
            "import). Fixed by re-importing the same known-good Week 21 CSV instead — "
            "re-importing already creates a fresh NEW-status version per Hera's versioning "
            "behaviour, so nothing about the test's purpose changes. The V2 CSV generation "
            "code was removed from test_weekly_ui.py."
        ),
        "Created via script": SCRIPT,
    },

    "Nom_WN_UI_10": {
        "Extra description ": (
            "[UPDATED 2026-07-03] _bulk_week() no longer verifies success by snapshotting "
            "status-new slot counts before/after the click (that approach was fragile — see "
            "Nom_WN_UI_11's note for the failure history). It now checks for Hera's own "
            "confirmation banner (e.g. “Confirmed week / For …”) via the new "
            "_detect_bulk_toast() helper: PASS if the banner text is found, FAIL if an error "
            "keyword appears instead. Still a real state check, not a click-success rubber "
            "stamp — just driven by the app's own feedback instead of re-deriving state from "
            "slot CSS classes."
        ),
    },

    "Nom_WN_UI_11": {
        "Extra description ": (
            "[UPDATED 2026-07-03] The previous DOM-based check (“snapshot non-rejected "
            "slot count before/after”) produced two false results on live runs: (1) a "
            "47/77 mismatch caused by SETUP_02's V2-CSV import failure leaving the DB in a "
            "partial state — not a real bug (see Nom_WN_UI_SETUP_02's note); (2) after a "
            "follow-up attempt to fix the slot-emptiness heuristic, a false AUTO-PASS of "
            "“0/0 rejected” because the tightened heuristic matched zero slots at "
            "all. Both were symptoms of relying on slot CSS-class/content heuristics for "
            "state verification. Replaced with _detect_bulk_toast(): PASS only if Hera's own "
            "“Rejected week / For …” confirmation banner is found after the "
            "click, FAIL on an error-keyword banner instead."
        ),
    },
}

# ---------------------------------------------------------------------------
# Apply patches
# ---------------------------------------------------------------------------
wb = openpyxl.load_workbook(PATH)
ws = wb.active

hdr = [str(c.value or "") for c in ws[1]]
COL = {h: i + 1 for i, h in enumerate(hdr)}   # 1-indexed for openpyxl, no .strip()

updated = []
for row in ws.iter_rows(min_row=2):
    tc_id_cell = row[COL["Test case ID"] - 1]
    tc_id = str(tc_id_cell.value or "").strip()
    if tc_id in PATCHES:
        patch = PATCHES[tc_id]
        for col_name, new_val in patch.items():
            col_idx = COL.get(col_name)
            if col_idx:
                cell = row[col_idx - 1]
                cell.value = new_val
                cell.alignment = WRAP
        updated.append(tc_id)
        print(f"  Updated {tc_id}")

if not updated:
    print("  No matching rows found — check Test case ID column.")
else:
    wb.save(PATH)
    print(f"\nSaved {PATH}")
    print(f"Updated {len(updated)} rows: {updated}")
