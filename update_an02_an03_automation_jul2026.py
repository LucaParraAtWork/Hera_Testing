"""
update_an02_an03_automation_jul2026.py
Credit the new AN_02/AN_03 negative-scenario automation added to
test_weekly_ui.py (2026-07-03) on the corresponding Excel rows and recolor
them from YELLOW (gap) to GREEN (automated).

  Nom_AN_02 -> test_weekly_ui.py AN_02 (_inconsistent_trailer_types):
    Drop-Off Type 1 saved on an earlier NEW slot, then a mismatched Pick-Up
    Type 2 attempted on a later NEW slot — expects rejection.
  Nom_AN_03 -> test_weekly_ui.py AN_03 (_two_dropoffs_without_pickup):
    a second trailer row forced to Drop-Off (no Pick-Up between) — expects
    rejection.

Both scenarios were written following the existing UI_02-05 modal-edit
pattern but have NOT been run live yet — selectors/dropdown behaviour may
need adjustment on the first real run (same caveat previously applied to
the Yearly Nomination script).

Run:
    python update_an02_an03_automation_jul2026.py
"""
import openpyxl
from openpyxl.styles import Alignment, PatternFill

PATH   = r"C:\Users\LucaParraGabarron\Desktop\ViryaH2\Hera_Testing\Hera_UAT_Test_Plan.xlsx"
SCRIPT = "test_weekly_ui.py"

WRAP  = Alignment(wrap_text=True, vertical="top")
GREEN = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")

# ---------------------------------------------------------------------------
# Patch definitions  { test_case_id : { column_name : new_value } }
# ---------------------------------------------------------------------------
PATCHES = {

    "Nom_AN_02": {
        "Extra description ": (
            "[AUTOMATED 2026-07-03] Covered by test_weekly_ui.py — AN_02 "
            "(_inconsistent_trailer_types): saves Drop-Off Type 1 on an earlier NEW "
            "slot, then attempts to save a mismatched Pick-Up Type 2 on a later NEW "
            "slot, and expects Hera to reject it (dropdown restriction, disabled Save, "
            "or an error message all count as a pass — expect=\"FAIL\" on the raw save "
            "outcome). NOT yet run live — dropdown/validation behaviour may need "
            "adjustment on the first real run."
        ),
        "Created via script": SCRIPT,
    },

    "Nom_AN_03": {
        "Extra description ": (
            "[AUTOMATED 2026-07-03] Covered by test_weekly_ui.py — AN_03 "
            "(_two_dropoffs_without_pickup): opens a NEW slot, sets Drop-Off Type 1, "
            "adds a second trailer row and forces it to Drop-Off too (no Pick-Up "
            "between), then attempts Save — expects Hera to reject it (dropdown "
            "restriction, disabled Save, or an error message all count as a pass — "
            "expect=\"FAIL\" on the raw save outcome). NOT yet run live — "
            "dropdown/validation behaviour may need adjustment on the first real run."
        ),
        "Created via script": SCRIPT,
    },
}

# ---------------------------------------------------------------------------
# Apply patches
# ---------------------------------------------------------------------------
wb = openpyxl.load_workbook(PATH)
ws = wb.active

hdr = [str(c.value or "") for c in ws[1]]
COL = {h: i + 1 for i, h in enumerate(hdr)}   # 1-indexed for openpyxl, no .strip()

updated = []
for row in ws.iter_rows(min_row=2):
    tc_id_cell = row[COL["Test case ID"] - 1]
    tc_id = str(tc_id_cell.value or "").strip()
    if tc_id in PATCHES:
        patch = PATCHES[tc_id]
        for col_name, new_val in patch.items():
            col_idx = COL.get(col_name)
            if col_idx:
                cell = row[col_idx - 1]
                cell.value = new_val
                cell.alignment = WRAP
        for cell in row:
            cell.fill = GREEN
        updated.append(tc_id)
        print(f"  Updated {tc_id} (recolored GREEN)")

if not updated:
    print("  No matching rows found — check Test case ID column.")
else:
    wb.save(PATH)
    print(f"\nSaved {PATH}")
    print(f"Updated {len(updated)} rows: {updated}")
