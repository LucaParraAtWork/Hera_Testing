"""
update_align_additional_nomination.py
Align the "Additional nomination" rows (Nom_AN_*) with the UI mechanics already
proven by test_weekly_ui.py, and record known coverage gaps.

Only touches: Nom_AN_01, Nom_AN_04, Nom_AN_06, Nom_AN_07.
  - Nom_AN_01 overlaps UI_12 (_add_slot_ui)          — also has a pre-existing
    legacy note pointing at hera_import_and_profiles.py::add_nomination_ui.
  - Nom_AN_04 overlaps UI_05 (_delete_trailer_row)    — gap: delete-to-empty edge case untested.
  - Nom_AN_06 overlaps UI_02 (_modify_offtaker)       — gap: CONFIRMED-slot case untested.
  - Nom_AN_07 overlaps UI_03 (_modify_trailer_type)   — gap: CONFIRMED-slot case untested.

Nom_AN_02, Nom_AN_03, Nom_AN_05 are intentionally left untouched: there is no
automation for them yet (inconsistent trailer types / two drops without
pick-up / delete slot entirely) — confirmed gap, not filled in.

Run:
    python update_align_additional_nomination.py
"""
import openpyxl
from openpyxl.styles import Alignment

PATH   = r"C:\Users\LucaParraGabarron\Desktop\ViryaH2\Hera_Testing\Hera_UAT_Test_Plan.xlsx"
SCRIPT = "test_weekly_ui.py"

WRAP = Alignment(wrap_text=True, vertical="top")

# ---------------------------------------------------------------------------
# Patch definitions  { test_case_id : { column_name : new_value } }
# Only "Created via script" and "Extra description " are touched — the
# hand-written title/steps/expected-results columns are left as-is.
# ---------------------------------------------------------------------------
PATCHES = {

    "Nom_AN_01": {
        "Extra description ": (
            "Automated: add_nomination_ui in hera_import_and_profiles.py (legacy). "
            "Also covered by test_weekly_ui.py – UI_12 (_add_slot_ui): click empty slot, "
            "select offtaker, add two trailer rows (Drop-Off + Pick-Up), Save. "
            "UI_12 is the current layered-script implementation; the "
            "hera_import_and_profiles.py version is kept for traceability only."
        ),
        "Created via script": SCRIPT,
    },

    "Nom_AN_04": {
        "Extra description ": (
            "Row-removal mechanic covered by test_weekly_ui.py – UI_05 (_delete_trailer_row), "
            "which requires a slot with >=2 trailer rows before deleting the last one. "
            "GAP: the 'slot updated or removed per business rules' edge case — deleting "
            "down to a single/zero transfer — is NOT exercised by UI_05."
        ),
        "Created via script": SCRIPT,
    },

    "Nom_AN_06": {
        "Extra description ": (
            "NEW-slot offtaker change covered by test_weekly_ui.py – UI_02 (_modify_offtaker). "
            "GAP: UI_02 only ever targets a NEW slot; the CONFIRMED-slot case named in this "
            "test ('confirmed or NEW slot') is not exercised."
        ),
        "Created via script": SCRIPT,
    },

    "Nom_AN_07": {
        "Extra description ": (
            "NEW-slot trailer-type change covered by test_weekly_ui.py – UI_03 "
            "(_modify_trailer_type). GAP: UI_03 only ever targets a NEW slot; a "
            "CONFIRMED-slot variant is not exercised."
        ),
        "Created via script": SCRIPT,
    },
}

# ---------------------------------------------------------------------------
# Apply patches
# ---------------------------------------------------------------------------
wb = openpyxl.load_workbook(PATH)
ws = wb.active

hdr = [str(c.value or "") for c in ws[1]]
COL = {h: i + 1 for i, h in enumerate(hdr)}   # 1-indexed for openpyxl

updated = []
for row in ws.iter_rows(min_row=2):
    tc_id_cell = row[COL["Test case ID"] - 1]
    tc_id = str(tc_id_cell.value or "").strip()
    if tc_id in PATCHES:
        patch = PATCHES[tc_id]
        for col_name, new_val in patch.items():
            col_idx = COL.get(col_name)
            if col_idx:
                cell = row[col_idx - 1]
                cell.value = new_val
                cell.alignment = WRAP
        updated.append(tc_id)
        print(f"  Updated {tc_id}")

if not updated:
    print("  No matching rows found — check Test case ID column.")
else:
    wb.save(PATH)
    print(f"\nSaved {PATH}")
    print(f"Updated {len(updated)} rows: {updated}")
