#!/usr/bin/env python3
"""
generate_dashboard.py – Business-facing UAT reporting dashboard
==================================================================
Reads a Hera UAT report (the master template, or — far more useful — one
of the timestamped copies under Reports/ that a live test run has already
written Pass/Fail/tester/comment into via uat_excel_reporter.py) and builds
a separate, clean, chart-driven Excel workbook meant to be shared with
non-technical stakeholders: at a glance, what's been tested and passed,
what's failed, and what hasn't been run yet — broken down by Process, then
Subprocess.

The input file is never modified. This always creates a brand new output
workbook.

Run:
    python generate_dashboard.py
        # picks the most recently modified file under Reports/ as input
        # (falls back to the master template if Reports/ is empty)

    python generate_dashboard.py --input "Reports/Hera_UAT_Test_Plan_20260910-153936.xlsx"

    python generate_dashboard.py --input ... --output "Dashboards/for_steering_committee.xlsx"
"""
from __future__ import annotations

import argparse
from collections import defaultdict
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import DataPoint
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

HERE = Path(__file__).resolve().parent
TEMPLATE_PATH = HERE / "Hera_UAT_Test_Plan.xlsx"
REPORTS_DIR = HERE / "Reports"
DASHBOARDS_DIR = HERE / "Dashboards"

ID_COL = "Test case ID"
PROCESS_COL = "Process"
SUBPROCESS_COL = "subprocess"
TITLE_COL = "Test case title"
STATUS_COL = "status"
TESTER_COL = "tester"
COMMENT_COL = "comment"
SCRIPT_COL = "Created via script"

# Friendlier label for each source sheet, used as a "coverage" column.
SHEET_LABELS = {
    "Covered by Scripts": "Automated",
    "Covered Manually": "Manual process",
    "Missing from Scripts": "Not yet automated",
}

BUCKET_ORDER = ["Passed", "Failed", "Error", "Inconclusive", "Skipped", "Not started"]
BUCKET_COLORS = {
    "Passed":       "2E7D32",  # green
    "Failed":       "C62828",  # red
    "Error":        "6A1B9A",  # purple
    "Inconclusive": "F9A825",  # amber
    "Skipped":      "757575",  # grey
    "Not started":  "B0BEC5",  # light grey-blue
}
HEADER_FILL = "1B5E20"   # dark green, echoes the Hera sidebar
HEADER_FONT_COLOR = "FFFFFF"


def _bucket(status) -> str:
    if status is None:
        return "Not started"
    s = str(status).strip()
    if not s or s.lower() == "not started":
        return "Not started"
    if s.upper() == "PASS":
        return "Passed"
    if s.upper() == "FAIL":
        return "Failed"
    if s.upper() == "ERROR":
        return "Error"
    if s == "?":
        return "Inconclusive"
    if s.upper() == "SKIP":
        return "Skipped"
    return "Inconclusive"


def _header_index(ws, name: str) -> int | None:
    for cell in ws[1]:
        if cell.value is not None and str(cell.value).strip() == name:
            return cell.column
    return None


def _pick_default_input() -> Path:
    if REPORTS_DIR.exists():
        candidates = [p for p in REPORTS_DIR.glob("*.xlsx") if not p.name.startswith("~$")]
        if candidates:
            return max(candidates, key=lambda p: p.stat().st_mtime)
    if not TEMPLATE_PATH.exists():
        raise FileNotFoundError(
            f"No report found under {REPORTS_DIR} and no master template at {TEMPLATE_PATH}.")
    return TEMPLATE_PATH


def load_rows(path: Path) -> list[dict]:
    """Read every test-case row across every sheet in the workbook into one flat list."""
    wb = openpyxl.load_workbook(path, data_only=True)
    rows: list[dict] = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        col = {
            "process": _header_index(ws, PROCESS_COL),
            "subprocess": _header_index(ws, SUBPROCESS_COL),
            "id": _header_index(ws, ID_COL),
            "title": _header_index(ws, TITLE_COL),
            "status": _header_index(ws, STATUS_COL),
            "tester": _header_index(ws, TESTER_COL),
            "comment": _header_index(ws, COMMENT_COL),
            "script": _header_index(ws, SCRIPT_COL),
        }
        if col["id"] is None:
            continue  # not a test-case sheet, skip
        for r in range(2, ws.max_row + 1):
            tc_id = ws.cell(row=r, column=col["id"]).value
            if not tc_id:
                continue

            def _get(key):
                c = col[key]
                return ws.cell(row=r, column=c).value if c else None

            status_raw = _get("status")
            rows.append({
                "process": (_get("process") or "").strip() if isinstance(_get("process"), str) else (_get("process") or "(no process)"),
                "subprocess": (_get("subprocess") or "").strip() if isinstance(_get("subprocess"), str) else (_get("subprocess") or "(no subprocess)"),
                "id": str(tc_id).strip(),
                "title": _get("title") or "",
                "status_raw": status_raw,
                "bucket": _bucket(status_raw),
                "tester": _get("tester") or "",
                "comment": _get("comment") or "",
                "coverage": SHEET_LABELS.get(sheet_name, sheet_name),
                "sheet": sheet_name,
            })
    return rows


def aggregate(rows: list[dict]):
    overall = {b: 0 for b in BUCKET_ORDER}
    by_process: dict[str, dict[str, int]] = defaultdict(lambda: {b: 0 for b in BUCKET_ORDER})
    by_process_sub: dict[tuple[str, str], dict[str, int]] = defaultdict(lambda: {b: 0 for b in BUCKET_ORDER})

    for r in rows:
        overall[r["bucket"]] += 1
        by_process[r["process"]][r["bucket"]] += 1
        by_process_sub[(r["process"], r["subprocess"])][r["bucket"]] += 1

    return overall, by_process, by_process_sub


# ---------------------------------------------------------------------------
# Styling helpers
# ---------------------------------------------------------------------------
THIN = Side(style="thin", color="D0D0D0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _style_header_row(ws, row: int, n_cols: int):
    for c in range(1, n_cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(bold=True, color=HEADER_FONT_COLOR)
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def _autosize(ws, widths: dict[int, int]):
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width


def _total(counts: dict[str, int]) -> int:
    return sum(counts.values())


def _pass_rate(counts: dict[str, int]) -> float:
    executed = counts["Passed"] + counts["Failed"] + counts["Error"] + counts["Inconclusive"] + counts["Skipped"]
    return (counts["Passed"] / executed) if executed else 0.0


# ---------------------------------------------------------------------------
# Sheet builders
# ---------------------------------------------------------------------------
# Charts live in their own column zone, well clear of every table, so their
# footprint (which openpyxl sizes in cm, not rows/columns) can never overlap
# the data next to them regardless of how many processes/buckets there are.
CHART_ANCHOR_COL = "O"
_CM_PER_ROW = 0.53   # ~15pt default row height
_ROW_GAP = 2


def _rows_for_cm(cm: float) -> int:
    return int(cm / _CM_PER_ROW) + 1


def _style_pie(pie: PieChart, bucket_order: list[str]):
    pie.height = 9
    pie.width = 16
    pie.title = "Overall status"
    pie.dataLabels = DataLabelList()
    pie.dataLabels.showPercent = True
    pie.dataLabels.showLegendKey = False
    pie.dataLabels.showSerName = False
    pie.dataLabels.showCatName = False
    pie.dataLabels.showVal = False
    pie.legend.position = "b"
    pie.legend.overlay = False
    series = pie.series[0]
    for i, bucket in enumerate(bucket_order):
        pt = DataPoint(idx=i)
        pt.graphicalProperties.solidFill = BUCKET_COLORS[bucket]
        series.data_points.append(pt)


def _style_bar(bar: BarChart, bucket_order: list[str]):
    bar.type = "col"
    bar.grouping = "stacked"
    bar.overlap = 100
    bar.gapWidth = 60
    bar.title = "Test status by process"
    bar.y_axis.title = "Test cases"
    bar.x_axis.delete = False
    bar.y_axis.delete = False
    bar.legend.position = "b"
    bar.legend.overlay = False
    for series, bucket in zip(bar.series, bucket_order):
        series.graphicalProperties.solidFill = BUCKET_COLORS[bucket]


def build_dashboard_sheet(wb, overall, by_process, source_path: Path):
    ws = wb.active
    ws.title = "Dashboard"

    ws["A1"] = "Hera UAT — Test Execution Dashboard"
    ws["A1"].font = Font(bold=True, size=18, color=HEADER_FILL)
    ws["A2"] = f"Source: {source_path.name}"
    ws["A2"].font = Font(italic=True, color="616161")
    ws["A3"] = f"Generated: {datetime.now():%Y-%m-%d %H:%M}"
    ws["A3"].font = Font(italic=True, color="616161")

    # ---- KPI strip (columns A-L) -----------------------------------------
    total = _total(overall)
    executed = total - overall["Not started"]
    kpis = [
        ("Total test cases", total, "424242"),
        ("Executed so far", f"{executed} ({executed / total:.0%} of total)" if total else "0",
         "1565C0"),
        ("Passed", overall["Passed"], BUCKET_COLORS["Passed"]),
        ("Failed", overall["Failed"] + overall["Error"], BUCKET_COLORS["Failed"]),
        ("Not started yet", overall["Not started"], BUCKET_COLORS["Not started"]),
        ("Pass rate — among tests executed so far", f"{_pass_rate(overall):.0%}", "1565C0"),
    ]
    start_row = 5
    ws.row_dimensions[start_row].height = 30
    for i, (label, value, color) in enumerate(kpis):
        col = 1 + i * 2
        c1 = ws.cell(row=start_row, column=col, value=label)
        c1.font = Font(bold=True, size=10, color="616161")
        c1.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c2 = ws.cell(row=start_row + 1, column=col, value=value)
        c2.font = Font(bold=True, size=20 if isinstance(value, int) else 14, color=color)
        c2.alignment = Alignment(horizontal="center")
        ws.merge_cells(start_row=start_row, start_column=col, end_row=start_row, end_column=col + 1)
        ws.merge_cells(start_row=start_row + 1, start_column=col, end_row=start_row + 1, end_column=col + 1)

    # ---- Overall status table (columns A-B; feeds the pie chart) -------
    table_row = start_row + 4
    ws.cell(row=table_row, column=1, value="Overall status").font = Font(bold=True, size=12)
    ws.cell(row=table_row + 1, column=1, value="Status")
    ws.cell(row=table_row + 1, column=2, value="Count")
    _style_header_row(ws, table_row + 1, 2)
    for i, bucket in enumerate(BUCKET_ORDER):
        r = table_row + 2 + i
        ws.cell(row=r, column=1, value=bucket)
        ws.cell(row=r, column=2, value=overall[bucket])
        fill = PatternFill("solid", fgColor=BUCKET_COLORS[bucket])
        ws.cell(row=r, column=1).fill = fill
        ws.cell(row=r, column=1).font = Font(color="FFFFFF", bold=True)
        ws.cell(row=r, column=2).border = BORDER
        ws.cell(row=r, column=1).border = BORDER

    pie = PieChart()
    data = Reference(ws, min_col=2, min_row=table_row + 1, max_row=table_row + 1 + len(BUCKET_ORDER))
    cats = Reference(ws, min_col=1, min_row=table_row + 2, max_row=table_row + 1 + len(BUCKET_ORDER))
    pie.add_data(data, titles_from_data=True)
    pie.set_categories(cats)
    _style_pie(pie, BUCKET_ORDER)
    pie_anchor_row = start_row
    ws.add_chart(pie, f"{CHART_ANCHOR_COL}{pie_anchor_row}")

    # ---- Per-process table (columns A-I; feeds the bar chart) ----------
    proc_table_row = table_row + 2 + len(BUCKET_ORDER) + 3
    ws.cell(row=proc_table_row, column=1, value="By process").font = Font(bold=True, size=12)
    header_row = proc_table_row + 1
    headers = ["Process", "Total"] + BUCKET_ORDER + ["Pass rate"]
    for c, h in enumerate(headers, start=1):
        ws.cell(row=header_row, column=c, value=h)
    _style_header_row(ws, header_row, len(headers))

    processes = sorted(by_process.keys())
    for i, proc in enumerate(processes):
        r = header_row + 1 + i
        counts = by_process[proc]
        ws.cell(row=r, column=1, value=proc).border = BORDER
        ws.cell(row=r, column=2, value=_total(counts)).border = BORDER
        for j, bucket in enumerate(BUCKET_ORDER):
            cell = ws.cell(row=r, column=3 + j, value=counts[bucket])
            cell.border = BORDER
            if counts[bucket]:
                cell.fill = PatternFill("solid", fgColor=BUCKET_COLORS[bucket])
                cell.font = Font(color="FFFFFF", bold=True)
        rate_cell = ws.cell(row=r, column=3 + len(BUCKET_ORDER), value=_pass_rate(counts))
        rate_cell.number_format = "0%"
        rate_cell.border = BORDER

    last_row = header_row + len(processes)
    bar = BarChart()
    cats = Reference(ws, min_col=1, min_row=header_row + 1, max_row=last_row)
    for j, bucket in enumerate(BUCKET_ORDER):
        data = Reference(ws, min_col=3 + j, min_row=header_row, max_row=last_row)
        bar.add_data(data, titles_from_data=True)
    bar.set_categories(cats)
    _style_bar(bar, BUCKET_ORDER)
    bar.height = 12
    bar.width = 22
    # Stack the bar chart below the pie chart, in the SAME chart column, with
    # a real gap computed from the pie's own height so they never touch.
    bar_anchor_row = pie_anchor_row + _rows_for_cm(pie.height) + _ROW_GAP
    ws.add_chart(bar, f"{CHART_ANCHOR_COL}{bar_anchor_row}")

    _autosize(ws, {1: 26, 2: 14, 3: 10, 4: 10, 5: 10, 6: 12, 7: 10, 8: 12, 9: 14, 10: 14, 11: 14, 12: 14})
    ws.column_dimensions["M"].width = 3  # narrow spacer column before the charts
    ws.column_dimensions["N"].width = 3
    ws.sheet_view.showGridLines = False


def build_process_summary_sheet(wb, by_process_sub):
    ws = wb.create_sheet("By Process & Subprocess")
    headers = ["Process", "Subprocess", "Total"] + BUCKET_ORDER + ["Pass rate"]
    for c, h in enumerate(headers, start=1):
        ws.cell(row=1, column=c, value=h)
    _style_header_row(ws, 1, len(headers))
    ws.freeze_panes = "A2"

    keys = sorted(by_process_sub.keys())
    row = 2
    current_process = None
    process_start_row = None
    grand = {b: 0 for b in BUCKET_ORDER}

    def _write_total_row(label, counts, bold_fill="ECEFF1"):
        nonlocal row
        ws.cell(row=row, column=1, value=label).font = Font(bold=True)
        ws.cell(row=row, column=2, value="")
        ws.cell(row=row, column=3, value=_total(counts)).font = Font(bold=True)
        for j, bucket in enumerate(BUCKET_ORDER):
            ws.cell(row=row, column=4 + j, value=counts[bucket]).font = Font(bold=True)
        rate_cell = ws.cell(row=row, column=4 + len(BUCKET_ORDER), value=_pass_rate(counts))
        rate_cell.number_format = "0%"
        rate_cell.font = Font(bold=True)
        for c in range(1, len(headers) + 1):
            ws.cell(row=row, column=c).fill = PatternFill("solid", fgColor=bold_fill)
            ws.cell(row=row, column=c).border = BORDER
        row += 1

    for process, subprocess in keys:
        if current_process is not None and process != current_process:
            proc_counts = {b: sum(by_process_sub[(current_process, sp)][b]
                                   for (p, sp) in keys if p == current_process) for b in BUCKET_ORDER}
            _write_total_row(f"{current_process} — Total", proc_counts)
            for r in range(process_start_row, row - 1):
                ws.row_dimensions[r].outline_level = 1
            ws.row_dimensions[row - 1].outline_level = 0
        if process != current_process:
            current_process = process
            process_start_row = row

        counts = by_process_sub[(process, subprocess)]
        for b in BUCKET_ORDER:
            grand[b] += counts[b]
        ws.cell(row=row, column=1, value=process).border = BORDER
        ws.cell(row=row, column=2, value=subprocess).border = BORDER
        ws.cell(row=row, column=3, value=_total(counts)).border = BORDER
        for j, bucket in enumerate(BUCKET_ORDER):
            cell = ws.cell(row=row, column=4 + j, value=counts[bucket])
            cell.border = BORDER
            if counts[bucket]:
                cell.fill = PatternFill("solid", fgColor=BUCKET_COLORS[bucket])
                cell.font = Font(color="FFFFFF", bold=True)
        rate_cell = ws.cell(row=row, column=4 + len(BUCKET_ORDER), value=_pass_rate(counts))
        rate_cell.number_format = "0%"
        rate_cell.border = BORDER
        row += 1

    if current_process is not None:
        proc_counts = {b: sum(by_process_sub[(current_process, sp)][b]
                               for (p, sp) in keys if p == current_process) for b in BUCKET_ORDER}
        _write_total_row(f"{current_process} — Total", proc_counts)
        for r in range(process_start_row, row - 1):
            ws.row_dimensions[r].outline_level = 1

    _write_total_row("GRAND TOTAL", grand, bold_fill="B0BEC5")

    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{row - 1}"
    ws.sheet_properties.outlinePr.summaryBelow = True
    _autosize(ws, {1: 26, 2: 30, 3: 8})
    for c in range(4, 4 + len(BUCKET_ORDER)):
        ws.column_dimensions[get_column_letter(c)].width = 11
    ws.column_dimensions[get_column_letter(4 + len(BUCKET_ORDER))].width = 11
    ws.sheet_view.showGridLines = False


def build_details_sheet(wb, rows: list[dict]):
    ws = wb.create_sheet("Details")
    headers = ["Process", "Subprocess", "Test case ID", "Title", "Coverage",
               "Status", "Tester", "Comment"]
    for c, h in enumerate(headers, start=1):
        ws.cell(row=1, column=c, value=h)
    _style_header_row(ws, 1, len(headers))
    ws.freeze_panes = "A2"

    rows_sorted = sorted(rows, key=lambda r: (r["process"], r["subprocess"], r["id"]))
    current_process = None
    group_start = 2
    for i, r in enumerate(rows_sorted, start=2):
        if current_process is not None and r["process"] != current_process:
            for rr in range(group_start, i):
                ws.row_dimensions[rr].outline_level = 1
            group_start = i
        current_process = r["process"]

        ws.cell(row=i, column=1, value=r["process"]).border = BORDER
        ws.cell(row=i, column=2, value=r["subprocess"]).border = BORDER
        ws.cell(row=i, column=3, value=r["id"]).border = BORDER
        title_cell = ws.cell(row=i, column=4, value=r["title"])
        title_cell.border = BORDER
        title_cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=i, column=5, value=r["coverage"]).border = BORDER
        status_cell = ws.cell(row=i, column=6, value=r["bucket"])
        status_cell.border = BORDER
        status_cell.fill = PatternFill("solid", fgColor=BUCKET_COLORS[r["bucket"]])
        status_cell.font = Font(color="FFFFFF", bold=True)
        status_cell.alignment = Alignment(horizontal="center")
        ws.cell(row=i, column=7, value=r["tester"]).border = BORDER
        comment_cell = ws.cell(row=i, column=8, value=r["comment"])
        comment_cell.border = BORDER
        comment_cell.alignment = Alignment(wrap_text=True, vertical="top")

    last_row = 1 + len(rows_sorted)
    if current_process is not None:
        for rr in range(group_start, last_row + 1):
            ws.row_dimensions[rr].outline_level = 1

    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{last_row}"
    ws.sheet_properties.outlinePr.summaryBelow = True
    _autosize(ws, {1: 22, 2: 26, 3: 16, 4: 45, 5: 16, 6: 12, 7: 16, 8: 40})
    ws.sheet_view.showGridLines = False


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------
def generate(input_path: Path, output_path: Path) -> None:
    rows = load_rows(input_path)
    if not rows:
        raise RuntimeError(f"No test-case rows found in {input_path}")
    overall, by_process, by_process_sub = aggregate(rows)

    wb = openpyxl.Workbook()
    build_dashboard_sheet(wb, overall, by_process, input_path)
    build_process_summary_sheet(wb, by_process_sub)
    build_details_sheet(wb, rows)
    wb.active = 0

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)

    total = _total(overall)
    print(f"\n  Dashboard written to: {output_path}")
    print(f"  Source               : {input_path}")
    print(f"  Total test cases     : {total}")
    for b in BUCKET_ORDER:
        print(f"    {b:<14}: {overall[b]}")
    print(f"  Pass rate (executed) : {_pass_rate(overall):.0%}\n")


def main():
    parser = argparse.ArgumentParser(
        description="Generate a business-facing UAT dashboard from a Hera UAT report.")
    parser.add_argument("--input", default=None, metavar="PATH",
                         help="Populated UAT report to summarize. Default: the most "
                              "recently modified file under Reports/, or the master "
                              "template if Reports/ is empty.")
    parser.add_argument("--output", default=None, metavar="PATH",
                         help="Where to write the dashboard. Default: a timestamped "
                              "file under Dashboards/.")
    args = parser.parse_args()

    input_path = Path(args.input) if args.input else _pick_default_input()
    if not input_path.exists():
        print(f"ERROR: input file not found: {input_path}")
        raise SystemExit(1)

    if args.output:
        output_path = Path(args.output)
    else:
        stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        output_path = DASHBOARDS_DIR / f"UAT_Dashboard_{stamp}.xlsx"

    generate(input_path, output_path)


if __name__ == "__main__":
    main()
