"""
update_sched_conf_rej_jul2026.py
Credits test_weekly_ui.py for the "Scheduler" process rows SCHED_CONF_01/02/03
and SCHED_REJ_01/02/03 on the main sheet.

User confirmed via a live run (2026-07-10) that these are the exact same
Confirm/Reject-button and Confirm-Week mechanism already covered by
Nom_WN_UI_06/07/08/09/10/11 — just filed under a different Process label
("Scheduler" vs "Weekly nomination – UI"). There is no separate Scheduler
screen/script in this repo. test_weekly_ui.py's SCENARIOS list was updated
in the same session to label UI_06/07/08/09/10/11 with their matching
SCHED_CONF_*/SCHED_REJ_* IDs.

Mapping:
  SCHED_CONF_01 <-> Nom_WN_UI_06 (UI_06: confirm single NEW slot)
  SCHED_REJ_01  <-> Nom_WN_UI_07 (UI_07: reject single confirmed slot)
  SCHED_CONF_02 <-> Nom_WN_UI_08 (UI_08: confirm all NEW slots one by one)
  SCHED_REJ_02  <-> Nom_WN_UI_09 (UI_09: reject all confirmed slots one by one)
  SCHED_CONF_03 <-> Nom_WN_UI_10 (UI_10: bulk Confirm Week)
  SCHED_REJ_03  <-> Nom_WN_UI_11 (UI_11: bulk Reject Week)

Run:
    python update_sched_conf_rej_jul2026.py
"""
import openpyxl
from openpyxl.styles import Alignment

PATH   = r"C:\Users\LucaParraGabarron\Desktop\ViryaH2\Hera_Testing\UAT Testing\Hera_UAT_Test_Plan.xlsx"
SCRIPT = "test_weekly_ui.py"

WRAP = Alignment(wrap_text=True, vertical="top")

PATCHES = {
    "SCHED_CONF_01": {
        "Extra description ": (
            "[AUDIT 2026-07-10, confirmed via live run] Same mechanism as "
            "Nom_WN_UI_06 (test_weekly_ui.py, UI_06) — confirm a single NEW slot "
            "via its modal's Confirm button. No separate Scheduler screen/script "
            "exists in this repo; this row and Nom_WN_UI_06 describe the same "
            "tested behaviour under a different Process label."
        ),
        "Created via script": SCRIPT,
    },
    "SCHED_REJ_01": {
        "Extra description ": (
            "[AUDIT 2026-07-10, confirmed via live run] Same mechanism as "
            "Nom_WN_UI_07 (test_weekly_ui.py, UI_07) — reject a single confirmed "
            "slot via its modal's Reject button."
        ),
        "Created via script": SCRIPT,
    },
    "SCHED_CONF_02": {
        "Extra description ": (
            "[AUDIT 2026-07-10, confirmed via live run] Same mechanism as "
            "Nom_WN_UI_08 (test_weekly_ui.py, UI_08) — iterate and confirm every "
            "NEW slot one by one."
        ),
        "Created via script": SCRIPT,
    },
    "SCHED_REJ_02": {
        "Extra description ": (
            "[AUDIT 2026-07-10, confirmed via live run] Same mechanism as "
            "Nom_WN_UI_09 (test_weekly_ui.py, UI_09) — iterate and reject every "
            "confirmed slot one by one."
        ),
        "Created via script": SCRIPT,
    },
    "SCHED_CONF_03": {
        "Extra description ": (
            "[AUDIT 2026-07-10, confirmed via live run] Same mechanism as "
            "Nom_WN_UI_10 (test_weekly_ui.py, UI_10) — bulk 'Confirm Week' via "
            "the Confirm Week button + Confirm popup choice, verified via Hera's "
            "own 'Confirmed week' banner."
        ),
        "Created via script": SCRIPT,
    },
    "SCHED_REJ_03": {
        "Extra description ": (
            "[AUDIT 2026-07-10, confirmed via live run] Same mechanism as "
            "Nom_WN_UI_11 (test_weekly_ui.py, UI_11) — bulk 'Reject Week' via "
            "the Confirm Week button + Reject popup choice, verified via Hera's "
            "own 'Rejected week' banner."
        ),
        "Created via script": SCRIPT,
    },
}

wb = openpyxl.load_workbook(PATH)
ws = wb.active

hdr = [str(c.value or "") for c in ws[1]]
COL = {h: i + 1 for i, h in enumerate(hdr)}   # 1-indexed for openpyxl, no .strip()

updated = []
for row in ws.iter_rows(min_row=2):
    tc_id_cell = row[COL["Test case ID"] - 1]
    tc_id = str(tc_id_cell.value or "").strip()
    if tc_id in PATCHES:
        patch = PATCHES[tc_id]
        for col_name, new_val in patch.items():
            col_idx = COL.get(col_name)
            if col_idx:
                cell = row[col_idx - 1]
                cell.value = new_val
                cell.alignment = WRAP
        updated.append(tc_id)
        print(f"  Updated {tc_id}")

if not updated:
    print("  No matching rows found — check Test case ID column.")
else:
    wb.save(PATH)
    print(f"\nSaved {PATH}")
    print(f"Updated {len(updated)} rows: {updated}")
