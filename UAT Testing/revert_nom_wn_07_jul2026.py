"""
revert_nom_wn_07_jul2026.py
Reverts the Nom_WN_07 credit from 2026-07-10: the NAV_01 scenario in
test_weekly_ui.py was removed (user watched it live and it was clicking
the wrong control — the date-picker calendar's exact markup was never
confirmed, only guessed at) — so the main sheet's "Created via script"
and "Extra description " for Nom_WN_07 are restored to their pre-NAV_01
state. Nom_WN_05 is untouched — only Nom_WN_07 is reverted.

Run:
    python revert_nom_wn_07_jul2026.py
"""
import openpyxl
from openpyxl.styles import Alignment

PATH = r"C:\Users\LucaParraGabarron\Desktop\ViryaH2\Hera_Testing\UAT Testing\Hera_UAT_Test_Plan.xlsx"

WRAP = Alignment(wrap_text=True, vertical="top")

PATCHES = {
    "Nom_WN_07": {
        "Extra description ": "Used in all automated steps.",
        "Created via script": None,
    },
}

wb = openpyxl.load_workbook(PATH)
ws = wb.active

hdr = [str(c.value or "") for c in ws[1]]
COL = {h: i + 1 for i, h in enumerate(hdr)}   # 1-indexed for openpyxl, no .strip()

updated = []
for row in ws.iter_rows(min_row=2):
    tc_id_cell = row[COL["Test case ID"] - 1]
    tc_id = str(tc_id_cell.value or "").strip()
    if tc_id in PATCHES:
        patch = PATCHES[tc_id]
        for col_name, new_val in patch.items():
            col_idx = COL.get(col_name)
            if col_idx:
                cell = row[col_idx - 1]
                cell.value = new_val
                cell.alignment = WRAP
        updated.append(tc_id)
        print(f"  Reverted {tc_id}")

if not updated:
    print("  No matching rows found — check Test case ID column.")
else:
    wb.save(PATH)
    print(f"\nSaved {PATH}")
    print(f"Reverted {len(updated)} rows: {updated}")
