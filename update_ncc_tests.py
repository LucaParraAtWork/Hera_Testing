"""
update_ncc_tests.py
Update the 4 NCC test cases covered by test_ncc_costs_breakdown.py.

Updates rows in-place: NCC_UI_01, NCC_UI_02, NCC_UI_05, NCC_CALC_07.
Fills in proper step-by-step input, corrected expected results, and the
'Created via script' column.  All other NCC rows are untouched.

Run:
    python update_ncc_tests.py
"""
import openpyxl
from openpyxl.styles import Alignment

PATH = r"C:\Users\LucaParraGabarron\Desktop\ViryaH2\Hera_Testing\Hera_UAT_Test_Plan.xlsx"
SCRIPT = "test_ncc_costs_breakdown.py"

WRAP = Alignment(wrap_text=True, vertical="top")

# ---------------------------------------------------------------------------
# Patch definitions  { test_case_id : { column_name : new_value } }
# ---------------------------------------------------------------------------
PATCHES = {

    "NCC_UI_01": {
        "Test case title": "Costs Breakdown page loads — 3 NCC Forecast cards visible",
        "subprocess": "Costs Breakdown UI",
        "Input with test step by step": (
            "1. Open Schedules > Costs Breakdown (/schedules/costs-breakdown).\n"
            "2. Wait for the page to reach network-idle (API: GetCostsBreakdownScreenAsync fires).\n"
            "3. Inspect the NCC Forecast section: verify exactly 3 cards are rendered.\n"
            "4. Verify each card shows an NCC value in €/kg.\n"
            "5. Screenshot saved automatically.\n"
            "\n"
            "Automated by: test_ncc_costs_breakdown.py (NCC_UI_01).\n"
            "Run: python \"NetCommodityCost/test_ncc_costs_breakdown.py\" --week N"
        ),
        "Expected results": (
            "3 NCC Forecast cards are visible: one per offtaker (Client A, Client B) and "
            "one Total Planning card. API returns netCommodityCostForecasts with 3 entries. "
            "Each entry has a numeric €/kg value."
        ),
        "Extra description ": (
            "Automated via test_ncc_costs_breakdown.py. "
            "The script navigates to /schedules/costs-breakdown, waits for the API response, "
            "and asserts len(netCommodityCostForecasts) == 3."
        ),
        "User": "Steering manager / Automation",
        "Created via script": SCRIPT,
    },

    "NCC_UI_02": {
        "Test case title": "7-day navigation cycle — every day triggers fresh API data",
        "subprocess": "Costs Breakdown UI",
        "Input with test step by step": (
            "1. Open Schedules > Costs Breakdown (/schedules/costs-breakdown).\n"
            "2. Navigate to the Monday of the target week using the date picker.\n"
            "3. Click the forward (→) arrow to step through Tuesday, Wednesday, …, Sunday.\n"
            "4. After each click verify GetCostsBreakdownScreenAsync fires and the NCC cards update.\n"
            "5. A screenshot is saved for each of the 7 days.\n"
            "\n"
            "Automated by: test_ncc_costs_breakdown.py (NCC_UI_02).\n"
            "Run: python \"NetCommodityCost/test_ncc_costs_breakdown.py\" --week N"
        ),
        "Expected results": (
            "API captured for all 7 days; 7/7 have NCC forecasts. "
            "NCC card values differ day-to-day according to scheduled activity. "
            "No day shows stale or missing data."
        ),
        "Extra description ": (
            "Automated via test_ncc_costs_breakdown.py. "
            "The script captures the CostsBreakdown POST response for each day by "
            "collecting all responses during navigation and keying each one by Brussels "
            "date (derived from the API’s ‘till’ timestamp)."
        ),
        "User": "Steering manager / Automation",
        "Created via script": SCRIPT,
    },

    "NCC_UI_05": {
        "Test case title": "€/MWh ↔ €/kg unit toggle — no errors, display updates",
        "subprocess": "Costs Breakdown UI",
        "Input with test step by step": (
            "1. Open Schedules > Costs Breakdown.\n"
            "2. Default unit is €/kg — verify NCC cards show values in €/kg.\n"
            "3. Click the unit toggle button to switch to €/MWh.\n"
            "4. Verify the NCC values rescale correctly (no ‘NaN’ or zero).\n"
            "5. Click the toggle again to switch back to €/kg.\n"
            "6. Verify no JS console error and no error toast appears throughout.\n"
            "\n"
            "Automated by: test_ncc_costs_breakdown.py (NCC_UI_05).\n"
            "Run: python \"NetCommodityCost/test_ncc_costs_breakdown.py\" --week N"
        ),
        "Expected results": (
            "€/kg → €/MWh toggle succeeds; no errors detected. "
            "Display updates with rescaled values. "
            "Toggling back to €/kg also works without errors."
        ),
        "Extra description ": (
            "Automated via test_ncc_costs_breakdown.py. "
            "The script clicks the toggle twice and asserts no page error state appears."
        ),
        "User": "Steering manager / Automation",
        "Created via script": SCRIPT,
    },

    "NCC_CALC_07": {
        "Test case title": (
            "Daily NCC Python vs API: |diff| < 0.02 €/kg or |pct| < 1 % — all 7 days"
        ),
        "subprocess": "Forecast € calculation validation",
        "Input with test step by step": (
            "Prerequisite: weekly nominations imported and confirmed for the target week.\n"
            "\n"
            "1. Capture g/s profiles for the target week:\n"
            "   • Navigate to Schedules > Profiles (ProfilesScreenAsync).\n"
            "   • Script reads truck-filling and gas-injection time-series (15-min slots).\n"
            "2. Navigate Costs Breakdown for all 7 days:\n"
            "   • Script collects every GetCostsBreakdownScreenAsync response that fires\n"
            "     during Mon→Sun navigation, keyed by Brussels calendar date.\n"
            "3. Load parameters from NetCommodityCost/parameters.csv:\n"
            "   • PRICE_ELEC (€/MWh), PRICE_GAS (€/MWh), ETA (electrolyser efficiency),\n"
            "     M_SOL (minimum stable operating level, MW).\n"
            "4. Run Python NCC reference calculation per day per offtaker:\n"
            "   NCC = (electricity cost for truck fillings\n"
            "          + gas MSOL cost\n"
            "          − gas injection revenue\n"
            "          − aFRR revenue)\n"
            "          / total H₂ kg\n"
            "5. For each day and each offtaker compare Python NCC vs API value.\n"
            "6. Assert tolerance: |diff| < 0.02 €/kg  OR  |pct_diff| < 1 %.\n"
            "\n"
            "Automated by: test_ncc_costs_breakdown.py (NCC_CALC_07).\n"
            "Run: python \"NetCommodityCost/test_ncc_costs_breakdown.py\" --week N\n"
            "Use --skip-capture to re-run calculation on existing artifacts.\n"
            "Prices can be overridden: --price-elec 85 --price-gas 28"
        ),
        "Expected results": (
            "21/21 comparisons pass (7 days × 3 entries: 2 offtakers + Total Planning). "
            "All absolute differences < 0.02 €/kg OR all percentage differences < 1 %. "
            "Terminal shows [OK] for every comparison; no [XX] entries."
        ),
        "Extra description ": (
            "Automated via test_ncc_costs_breakdown.py. "
            "The Python NCC calculation ports the reference formula from NCC_TESTING.py "
            "(English, streamlined). "
            "Trailer-to-offtaker allocation uses confirmed pickups from the scheduling "
            "artifact (closest pickup within ±2 h window); falls back to 50/50 split "
            "if no scheduling artifact is available for the week."
        ),
        "User": "Steering manager / Automation",
        "Created via script": SCRIPT,
    },
}

# ---------------------------------------------------------------------------
# Apply patches
# ---------------------------------------------------------------------------
wb = openpyxl.load_workbook(PATH)
ws = wb.active

hdr = [str(c.value or "").strip() for c in ws[1]]
COL = {h: i + 1 for i, h in enumerate(hdr)}   # 1-indexed for openpyxl

updated = []
for row in ws.iter_rows(min_row=2):
    tc_id_cell = row[COL["Test case ID"] - 1]
    tc_id = str(tc_id_cell.value or "").strip()
    if tc_id in PATCHES:
        patch = PATCHES[tc_id]
        for col_name, new_val in patch.items():
            col_idx = COL.get(col_name)
            if col_idx:
                cell = row[col_idx - 1]
                cell.value = new_val
                cell.alignment = WRAP
        updated.append(tc_id)
        print(f"  Updated {tc_id}")

if not updated:
    print("  No matching rows found — check Test case ID column.")
else:
    wb.save(PATH)
    print(f"\nSaved {PATH}")
    print(f"Updated {len(updated)} rows: {updated}")
