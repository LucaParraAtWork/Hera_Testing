"""
create_script_coverage_sheets_jul2026.py
Adds two new sheets to Hera_UAT_Test_Plan.xlsx, built from a line-by-line
code audit (not just the "Created via script" column) of all 9 test
scripts across the 8 folders that run_all_tests.py drives:

  Weekly Nomination, Yearly Nomination, Master Data, Planned Maintenance,
  Scheduled Transfers, NetCommodityCost, EndToEndPlannedMaintenance,
  Plant Communication - sending (no test script there).

Sheet 1 "Covered by Scripts": every row that a script GENUINELY,
concretely implements (real assertion tied to that exact ID), in the
order run_all_tests.py executes the folders/scripts, and within each
script in its own scenario execution order — same 12 columns as the main
sheet.

Sheet 2 "Missing from Scripts": every other row from the main sheet
(everything not in Sheet 1), in the main sheet's original order.

This audit deliberately does NOT just trust the main sheet's existing
"Created via script" column — several discrepancies were found and are
called out via a short appended note in "Extra description " on the
affected row, ONLY within these two new sheets (the main sheet is left
completely untouched):

  - Nom_YN_05/06/07: test_yearly_ui.py genuinely implements these (real
    API-response assertions) but the main sheet deliberately leaves
    "Created via script" blank pending a live run (per prior explicit
    decision) — included in Sheet 1 with a note, main sheet unchanged.
  - ACS_MD_10: main sheet currently credits test_master_data.py, but
    _run_ACS_MD_10() unconditionally returns PASS regardless of state
    (no real assertion) — EXCLUDED from Sheet 1, placed in Sheet 2 with
    a note explaining the fake-pass finding.
  - Nom_WN_08: real scenario exists (shared with Nom_WN_01) but only
    checks generic import-success keywords, not the specific "New status"
    claim in its title — kept in Sheet 1 with a caveat note.
  - PROF_02/OCC_02: only indirectly covered (validated as a precondition
    by other Scheduled Transfers scenarios, no dedicated confirm action)
    — kept in Sheet 1 with a caveat note.
  - SCHED_TT_01, NCC_UI_06: possible duplicates of other rows — left in
    Sheet 2 with a note for a human decision.
  - SCHED_CONF_01/02/03, SCHED_REJ_01/02/03: user confirmed via a live
    run (2026-07-10) these are the same Confirm/Reject-button and
    Confirm-Week mechanism as Nom_WN_UI_06/07/08/09/10/11 — moved to
    Sheet 1 (test_weekly_ui.py), main sheet credited too via
    update_sched_conf_rej_jul2026.py, and the script's own SCENARIOS
    list now labels UI_06/07/08/09/10/11 with their matching SCHED_*
    IDs so the mapping is explicit in the code, not just the Excel.

Run:
    python create_script_coverage_sheets_jul2026.py
"""
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

PATH = r"C:\Users\LucaParraGabarron\Desktop\ViryaH2\Hera_Testing\Hera_UAT_Test_Plan.xlsx"
MAIN_SHEET = "Functional testing | UAT"

WRAP = Alignment(wrap_text=True, vertical="top")
HEADER_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
HEADER_FONT = Font(bold=True)

# ---------------------------------------------------------------------------
# Sheet 1 — ordered list of (Test case ID, optional note-to-append) exactly
# as verified against the code, in run_all_tests.py folder order and each
# script's own scenario execution order. AN_01/04/06/07 are placed right
# after the UI_XX scenario they share code with (not alphabetically) to
# reflect true execution grouping.
# ---------------------------------------------------------------------------
COVERED = [
    # === Weekly Nomination ===============================================
    # test_weekly_import.py (TEST_CASES order)
    ("Nom_WN_01", None),
    ("Nom_WN_08", "[AUDIT 2026-07-03] Real scenario exists (shared TEST_CASES entry "
                  "with Nom_WN_01), but the automated check only scans for generic "
                  "import-success keywords — it does not assert the specific 'New "
                  "status' DOM claim in this row's title. Kept as covered but weaker "
                  "than the other credited rows."),
    ("Nom_WN_03", None),
    ("Nom_WN_04", None),
    ("Nom_WN_05", "[AUDIT 2026-07-10] New TEST_CASES entry: imports "
                  "Messer_Nomination_Week26.csv, a fixed past week (26/2026), "
                  "using the stricter _detect_success_toast_only() detector "
                  "(FAIL unless an explicit success toast is found). Not yet "
                  "run live."),
    ("Nom_WN_NEG_03", None),
    ("Nom_WN_NEG_05", None),
    ("Nom_WN_NEG_06", None),
    ("Nom_WN_NEG_07", None),
    ("Nom_WN_NEG_08", None),
    ("Nom_WN_NEG_10", None),
    ("Nom_WN_NEG_13", None),
    ("Nom_WN_NEG_15", None),
    ("Nom_WN_NEG_16", "[AUDIT 2026-07-03] Real, ID-matching scenario exists, but its "
                      "own expect=\"?\" — purely observational/human-judgment, no "
                      "programmatic pass/fail."),
    ("Nom_WN_NEG_17", None),
    ("Nom_WN_NEG_18", None),
    # test_weekly_ui.py (SCENARIOS n order; AN_* overlaps placed next to their shared code)
    ("Nom_WN_UI_SETUP_01", None),
    ("Nom_WN_UI_01", None),
    ("Nom_WN_UI_02", None),
    ("Nom_AN_06", "[AUDIT 2026-07-03] Shares code with Nom_WN_UI_02 (_modify_offtaker) "
                  "— genuine for the NEW-slot case; the CONFIRMED-slot variant named "
                  "in this test's own steps is not exercised (documented gap)."),
    ("Nom_WN_UI_03", None),
    ("Nom_AN_07", "[AUDIT 2026-07-03] Shares code with Nom_WN_UI_03 (_modify_trailer_type) "
                  "— genuine for the NEW-slot case; the CONFIRMED-slot variant is not "
                  "exercised (documented gap)."),
    ("Nom_WN_UI_04", None),
    ("Nom_WN_UI_05", None),
    ("Nom_AN_04", "[AUDIT 2026-07-03] Shares code with Nom_WN_UI_05 (_delete_trailer_row) "
                  "— genuine for the >=2-row case; deleting down to a single/zero "
                  "transfer is Nom_AN_05's job, not this row's (documented gap)."),
    ("Nom_AN_02", "[AUDIT 2026-07-03] Real scenario (_inconsistent_trailer_types), "
                  "written 2026-07-03 — not yet run live, dropdown/validation "
                  "behaviour may need adjustment on first real run."),
    ("Nom_AN_03", "[AUDIT 2026-07-03] Real scenario (_two_dropoffs_without_pickup), "
                  "written 2026-07-03 — not yet run live, same caveat as Nom_AN_02."),
    ("Nom_AN_05", "[AUDIT 2026-07-03] Real scenario (_delete_slot_entirely), written "
                  "2026-07-03 — not yet run live, same caveat as Nom_AN_02/03."),
    ("Nom_WN_UI_06", None),
    ("SCHED_CONF_01", "[AUDIT 2026-07-10, confirmed via live run] Same mechanism as "
                      "Nom_WN_UI_06 (test_weekly_ui.py, UI_06) — confirm a single NEW "
                      "slot. No separate Scheduler screen/script exists in this repo."),
    ("Nom_WN_UI_07", None),
    ("SCHED_REJ_01", "[AUDIT 2026-07-10, confirmed via live run] Same mechanism as "
                     "Nom_WN_UI_07 (test_weekly_ui.py, UI_07) — reject a single "
                     "confirmed slot."),
    ("Nom_WN_UI_08", None),
    ("SCHED_CONF_02", "[AUDIT 2026-07-10, confirmed via live run] Same mechanism as "
                      "Nom_WN_UI_08 (test_weekly_ui.py, UI_08) — confirm all NEW "
                      "slots one by one."),
    ("Nom_WN_UI_09", None),
    ("SCHED_REJ_02", "[AUDIT 2026-07-10, confirmed via live run] Same mechanism as "
                     "Nom_WN_UI_09 (test_weekly_ui.py, UI_09) — reject all confirmed "
                     "slots one by one."),
    ("Nom_WN_UI_SETUP_02", None),
    ("Nom_WN_UI_10", None),
    ("SCHED_CONF_03", "[AUDIT 2026-07-10, confirmed via live run] Same mechanism as "
                      "Nom_WN_UI_10 (test_weekly_ui.py, UI_10) — bulk Confirm Week, "
                      "verified via Hera's own 'Confirmed week' banner."),
    ("Nom_WN_UI_11", None),
    ("SCHED_REJ_03", "[AUDIT 2026-07-10, confirmed via live run] Same mechanism as "
                     "Nom_WN_UI_11 (test_weekly_ui.py, UI_11) — bulk Reject Week, "
                     "verified via Hera's own 'Rejected week' banner."),
    ("Nom_WN_UI_12", None),
    ("Nom_AN_01", "[AUDIT 2026-07-03] Shares code with Nom_WN_UI_12 (_add_slot_ui) — "
                  "unambiguous match to this row's 'happy path' title."),
    ("Nom_WN_07", "[AUDIT 2026-07-10] New NAV_01 scenario: opens the date-picker "
                  "calendar via the current-week label (not </>), picks a day "
                  "from an adjacent month, verifies the week updates, then "
                  "returns to the target week and verifies no stale data. Not "
                  "yet run live — calendar selectors are a best-effort guess."),

    # === Yearly Nomination ================================================
    # test_yearly_import.py (TEST_CASES order)
    ("Nom_YN_01", None),
    ("Nom_YN_03", None),
    ("Nom_YN_04", None),
    ("Nom_YN_NEG_03", None),
    ("Nom_YN_NEG_04", None),
    ("Nom_YN_NEG_05", None),
    ("Nom_YN_NEG_06", None),
    ("Nom_YN_NEG_07", None),
    ("Nom_YN_NEG_08", None),
    ("Nom_YN_NEG_09", None),
    ("Nom_YN_NEG_10", None),
    ("Nom_YN_NEG_12", None),
    ("Nom_YN_NEG_14", None),
    ("Nom_YN_NEG_15", None),
    # test_yearly_ui.py (SCENARIOS n order: UI_01/02/03 -> Excel 05/07/06)
    ("Nom_YN_05", "[AUDIT 2026-07-03] Confirmed genuinely implemented (UI_01: asserts "
                  "currentVersion/versionLines/breakdowns are non-empty from the live "
                  "GetYearlyNominationScreenAsync response). Main sheet's 'Created via "
                  "script' is deliberately left blank pending a live run of this "
                  "script — that policy is unchanged; this sheet reflects code content."),
    ("Nom_YN_07", "[AUDIT 2026-07-03] Confirmed genuinely implemented (UI_02: chart "
                  "bounding-box check + cross-validates monthly/weekly breakdowns "
                  "against the baseline CSV). Same 'not yet credited on main sheet, "
                  "pending live run' note as Nom_YN_05."),
    ("Nom_YN_06", "[AUDIT 2026-07-03] Confirmed genuinely implemented (UI_03: edits a "
                  "card value, submits, asserts new version = old+1, status='new', "
                  "prior version still in history). Same 'not yet credited on main "
                  "sheet, pending live run' note as Nom_YN_05."),

    # === Master Data =======================================================
    ("ACS_MD_01", None),
    ("ACS_MD_02", None),
    ("ACS_TM_CO_01", None),
    ("ACS_TM_CO_02", None),
    ("ACS_TM_CO_03", None),
    ("ACS_MD_03", None),
    ("ACS_MD_04", None),
    ("ACS_TM_QUAL_01", None),
    ("ACS_TM_QUAL_04", None),
    ("ACS_TM_TRK_01", None),
    ("ACS_TM_TRK_02", None),
    ("ACS_TM_TRK_03", None),
    ("ACS_TM_TRK_04", None),
    ("ACS_TM_TRK_05", None),
    ("ACS_TM_TRK_06", None),
    ("ACS_MD_05", None),
    ("ACS_MD_06", None),
    ("ACS_TM_QUAL_03", None),
    ("ACS_TM_QUAL_05", None),
    ("ACS_TM_DRV_01", None),
    ("ACS_TM_DRV_02", None),
    ("ACS_TM_DRV_03", None),
    ("ACS_TM_DRV_04", None),
    ("ACS_MD_07", None),
    ("ACS_MD_08", None),
    ("ACS_TM_QUAL_02", None),
    ("ACS_TM_TRL_01", None),
    ("ACS_TM_TRL_02", None),
    ("ACS_TM_TRL_03", None),
    ("ACS_MD_09", "[AUDIT 2026-07-03] Genuine but shallow: PASS iff row count > 1, "
                  "no assertion on the actual row content/data."),
    # NOTE: ACS_MD_10 intentionally excluded — see Sheet 2.

    # === Planned Maintenance ===============================================
    ("MAINT_CR_01", None),
    ("MAINT_CR_02", None),
    ("MAINT_CR_03", None),
    ("MAINT_ED_01", "[AUDIT 2026-07-03] Real persistence assertion (re-opens and "
                    "reads back the changed value), but semi-manual — requires a "
                    "human to click the orange Gantt block (canvas-based, no "
                    "reliable selector)."),
    ("MAINT_ED_02", "[AUDIT 2026-07-03] Same semi-manual caveat as MAINT_ED_01."),
    ("MAINT_ED_03", None),
    ("MAINT_DEL_01", "[AUDIT 2026-07-03] Real API-based assertion (block absent from "
                     "GetMaintenanceScreenAsync afterwards), semi-manual click, and "
                     "falls back to an automatic PASS if the API listener doesn't "
                     "fire in time — a weak fallback path worth knowing about."),
    ("MAINT_PR_01", None),
    ("MAINT_PR_02", None),
    ("MAINT_PR_03", None),
    ("MAINT_PR_04", None),

    # === Scheduled Transfers ===============================================
    ("SCHED_TRF_01", None),
    ("SCHED_TRF_02", "[AUDIT 2026-07-03] Own code comment marks the segment/session "
                     "table as 'reference display (no checks here)' — the actual "
                     "count comparison implied by this row's title happens in "
                     "SCHED_TRF_03, not here."),
    ("SCHED_TRF_03", None),
    ("SCHED_TRF_04", None),
    ("SCHED_TRF_05", None),
    ("SCHED_TRF_06", None),
    ("SCHED_TRF_07", None),
    ("SCHED_TRF_08", None),
    ("PROF_02", "[AUDIT 2026-07-03] Indirect only: no dedicated 'confirm then check' "
               "scenario exists — SCHED_TRF_03/04/05/07 validate Profiles content "
               "that assumes confirmed data as a precondition, rather than "
               "performing the confirm action and observing the transition."),
    ("OCC_02", "[AUDIT 2026-07-03] Same indirect-only caveat as PROF_02, via "
              "SCHED_TRF_05/06/07's occupancy-count/derivation checks."),
    ("PROF_03", "[AUDIT 2026-07-03] Covered by SCHED_TRF_08's explicit "
               "post-reject-all assertion (segment count == 0)."),
    ("OCC_03", "[AUDIT 2026-07-03] Covered by SCHED_TRF_08's explicit "
              "post-reject-all assertion (non-maintenance occupation count == 0)."),

    # === NetCommodityCost ===================================================
    ("NCC_UI_01", None),
    ("NCC_UI_02", None),
    ("NCC_UI_05", None),
    ("NCC_CALC_07", None),

    # === EndToEndPlannedMaintenance =========================================
    ("E2E_MAINT_01", None),
    ("E2E_MAINT_02", None),

    # Plant Communication - sending: no test script, nothing to add.
]

# ---------------------------------------------------------------------------
# IDs where the main sheet's "Created via script" is blank on purpose (policy
# decisions unrelated to whether the code exists) — filled in ONLY on this
# new sheet, since its entire point is "what's genuinely in the scripts."
# ---------------------------------------------------------------------------
FILL_CREATED_VIA_SCRIPT = {
    "Nom_YN_05": "test_yearly_ui.py",
    "Nom_YN_06": "test_yearly_ui.py",
    "Nom_YN_07": "test_yearly_ui.py",
}

# ---------------------------------------------------------------------------
# Extra notes to append on specific Sheet-2 (gap) rows, explaining findings
# from this audit that go beyond "just uncredited."
# ---------------------------------------------------------------------------
GAP_NOTES = {
    "ACS_MD_10": (
        "[AUDIT 2026-07-03] EXCLUDED from 'Covered by Scripts' despite main sheet "
        "crediting test_master_data.py: _run_ACS_MD_10() unconditionally returns "
        "PASS regardless of page state (counts rows for a descriptive string only, "
        "no conditional) — confirmed fake-pass, not a real check."
    ),
    "SCHED_TT_01": (
        "[AUDIT 2026-07-03] Possibly the same underlying behaviour as "
        "SCHED_TRF_06 (Scheduled Transfers) — that scenario derives Drop/Filling/"
        "Pick windows from Drop-Off/Pick-Up slot times, which sounds conceptually "
        "identical to this row's title. Not marked covered because the overlap "
        "isn't certain — needs a human read of this row's exact step text vs. "
        "SCHED_TRF_06's code."
    ),
    "NCC_UI_06": (
        "[AUDIT 2026-07-03] Possible stale duplicate of NCC_UI_01 (flagged in the "
        "2026-07-03 alignment review too) — different subprocess grouping "
        "('Forecast' vs 'Costs Breakdown UI') means it might be a distinct, "
        "broader requirement instead. Not covered by test_ncc_costs_breakdown.py "
        "either way — needs a human call on whether it's a duplicate."
    ),
}


def main():
    wb = openpyxl.load_workbook(PATH)
    main_ws = wb[MAIN_SHEET]

    hdr = [str(c.value) for c in main_ws[1]]
    col_idx = {h.strip(): i for i, h in enumerate(hdr)}  # 0-indexed, stripped for lookup
    extra_col_name = next(h for h in hdr if h.strip() == "Extra description")

    by_id = {}
    for row in main_ws.iter_rows(min_row=2, values_only=False):
        tcid_cell = row[col_idx["Test case ID"]]
        tcid = str(tcid_cell.value).strip() if tcid_cell.value else None
        if tcid:
            by_id[tcid] = [c.value for c in row]

    covered_ids_in_order = [tcid for tcid, _ in COVERED]
    covered_set = set(covered_ids_in_order)

    missing_ids = [row[col_idx["Test case ID"]] for row in
                   ([c.value for c in r] for r in main_ws.iter_rows(min_row=2))
                   if row[col_idx["Test case ID"]] and
                   str(row[col_idx["Test case ID"]]).strip() not in covered_set]
    # de-dupe while preserving order (shouldn't be needed, IDs are unique, but safe)
    seen = set()
    missing_ids_ordered = []
    for m in missing_ids:
        m = str(m).strip()
        if m not in seen:
            seen.add(m)
            missing_ids_ordered.append(m)

    def _write_sheet(name, ids_with_notes, fill_created_via=None):
        if name in wb.sheetnames:
            del wb[name]
        ws = wb.create_sheet(name)
        for i, h in enumerate(hdr, 1):
            cell = ws.cell(row=1, column=i, value=h)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
        for i, h in enumerate(hdr, 1):
            ws.column_dimensions[get_column_letter(i)].width = \
                main_ws.column_dimensions[get_column_letter(i)].width

        r = 2
        for tcid, note in ids_with_notes:
            if tcid not in by_id:
                print(f"  [WARN] '{tcid}' not found in main sheet — skipped")
                continue
            values = list(by_id[tcid])
            if note:
                extra_idx = col_idx["Extra description"]
                existing = values[extra_idx] or ""
                values[extra_idx] = (existing + "\n" + note).strip() if existing else note
            if fill_created_via and tcid in fill_created_via:
                cvs_idx = col_idx["Created via script"]
                if not values[cvs_idx]:
                    values[cvs_idx] = fill_created_via[tcid]
            for c, v in enumerate(values, 1):
                cell = ws.cell(row=r, column=c, value=v)
                cell.alignment = WRAP
            r += 1
        print(f"  Wrote {r - 2} rows to '{name}'")

    _write_sheet("Covered by Scripts", COVERED, fill_created_via=FILL_CREATED_VIA_SCRIPT)

    missing_with_notes = [(tcid, GAP_NOTES.get(tcid)) for tcid in missing_ids_ordered]
    _write_sheet("Missing from Scripts", missing_with_notes)

    wb.save(PATH)
    print(f"\nSaved {PATH}")
    print(f"Covered by Scripts   : {len(COVERED)} rows")
    print(f"Missing from Scripts : {len(missing_ids_ordered)} rows")
    print(f"Total (should be 206): {len(COVERED) + len(missing_ids_ordered)}")


if __name__ == "__main__":
    main()
